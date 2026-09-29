"""Build the TalkNFace LCT 2026 pitch from the official template canvases."""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path
from typing import Optional

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

INK = RGBColor(0x1C, 0x1D, 0x22)
PURPLE = RGBColor(0x31, 0x0F, 0x53)
PINK = RGBColor(0xFC, 0x37, 0x77)

ROOT = Path(__file__).resolve().parents[3]
ASSETS = ROOT / "docs" / "submission" / "assets"
EVIDENCE = ROOT / "docs" / "gates" / "evidence"
PPTX = ROOT / "docs" / "submission" / "TALKNFACE_PRODUCT_PITCH_RU.pptx"
TMP = ROOT / ".tmp" / "organizer-presentation-01"
ASSEMBLED = TMP / "assembled.pptx"
ASSEMBLE_PS1 = Path(__file__).resolve().parent / "assemble_template.ps1"
TEAM_LOGO_SRC = ROOT / "photo_2026-09-29_23-15-40.jpg"
TEAM_LOGO = ASSETS / "khodoki-logo.jpg"
BLOB = "https://github.com/AiratBastanov/TalkNFace/blob/docs/organizer-presentation-01/"


def find_template() -> Path:
    names = list(ROOT.glob("ЛЦТ2026*.pptx"))
    if names:
        return names[0]
    desktop = Path.home() / "Desktop"
    found = list(desktop.glob("ЛЦТ2026*.pptx")) if desktop.exists() else []
    if found:
        return found[0]
    raise FileNotFoundError("Official LCT 2026 template PPTX not found next to the repo or on Desktop.")


def crop(src: Path, box: tuple[int, int, int, int], dest: Path) -> None:
    im = Image.open(src).convert("RGB")
    cut = im.crop(box)
    framed = Image.new("RGB", (cut.width + 2, cut.height + 2), (220, 229, 226))
    framed.paste(cut, (1, 1))
    dest.parent.mkdir(parents=True, exist_ok=True)
    framed.save(dest, "PNG", optimize=True)


def prepare_assets() -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    files = {
        "hero": (EVIDENCE / "app-qwen-independent-01" / "app-s2-result-1280.png", (150, 78, 1130, 1000)),
        "settings": (EVIDENCE / "app-g5-context-configuration-01" / "g5-initial-1280-viewport.png", (40, 240, 1240, 820)),
        "evidence": (EVIDENCE / "app-g6-deterministic-feedback-01" / "g6-s1-evidence-390.png", (4, 0, 386, 530)),
        "process100": (EVIDENCE / "app-clean-machine-reproducibility-01" / "clean-feedback-S1.png", (8, 88, 382, 800)),
    }
    out = {}
    for name, (src, box) in files.items():
        dest = ASSETS / f"{name}.png"
        crop(src, box, dest)
        out[name] = dest
    if not TEAM_LOGO_SRC.exists():
        raise FileNotFoundError(f"Team logo not found: {TEAM_LOGO_SRC}")
    shutil.copyfile(TEAM_LOGO_SRC, TEAM_LOGO)
    out["khodoki"] = TEAM_LOGO
    return out


def extract_alabuga(template: Path) -> Path:
    """Copy the official Алабуга mark from template slide 6; knock out the black sheet."""
    prs = Presentation(str(template))
    slide = prs.slides[5]
    dest = TMP / "alabuga-logo.png"
    dest.parent.mkdir(parents=True, exist_ok=True)
    picked = None
    for sh in slide.shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        if int(sh.left) == 7956509:
            picked = sh
            break
    if picked is None:
        raise RuntimeError("Could not find the Алабуга picture on template slide 6.")
    raw = TMP / "alabuga-raw.png"
    raw.write_bytes(picked.image.blob)
    im = Image.open(raw).convert("RGBA")
    pixels = []
    for r, g, b, a in im.getdata():
        if r < 40 and g < 40 and b < 40:
            pixels.append((255, 255, 255, 0))
        else:
            pixels.append((r, g, b, a))
    im.putdata(pixels)
    im.save(dest, "PNG")
    return dest


def put(shape, text: str, color: Optional[RGBColor] = None) -> None:
    """Replace text, drop leftover empty bullets, keep the theme font."""
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    tf = shape.text_frame
    tf.word_wrap = True
    lines = [ln for ln in text.split("\n") if ln.strip() != ""]
    if not lines:
        lines = [""]
    paras = list(tf.paragraphs)
    for i, line in enumerate(lines):
        if i < len(paras):
            p = paras[i]
        else:
            p = tf.add_paragraph()
            p.alignment = paras[0].alignment
        if p.runs:
            p.runs[0].text = line
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.add_run().text = line
        if color is not None:
            for run in p.runs:
                if run.text:
                    run.font.color.rgb = color
    body = tf._txBody
    xml_paras = body.findall(qn("a:p"))
    for el in xml_paras[len(lines) :]:
        body.remove(el)


def shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def drop_shape(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def insert_into_placeholder(slide, image: Path, name: str) -> None:
    ph = shape_by_name(slide, name)
    ph.insert_picture(str(image))


def add_picture_fit(slide, image: Path, left, top, width, height) -> None:
    im = Image.open(image)
    aspect = im.width / float(im.height)
    box_aspect = width / float(height)
    if aspect > box_aspect:
        h = int(width / aspect)
        t = top + (height - h) // 2
        slide.shapes.add_picture(str(image), left, t, width, h)
    else:
        w = int(height * aspect)
        l = left + (width - w) // 2
        slide.shapes.add_picture(str(image), l, top, w, height)


def strip_notes(prs: Presentation) -> None:
    for slide in prs.slides:
        if not slide.has_notes_slide:
            continue
        tf = slide.notes_slide.notes_text_frame
        tf.text = ""


def assemble(template: Path) -> None:
    TMP.mkdir(parents=True, exist_ok=True)
    cmd = [
        "powershell",
        "-NoProfile",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(ASSEMBLE_PS1),
        "-Template",
        str(template),
        "-OutFile",
        str(ASSEMBLED),
    ]
    subprocess.run(cmd, check=True)


def fill_title(slide, logo: Path) -> None:
    put(shape_by_name(slide, "Заголовок 2"), "Ходоки")
    put(
        shape_by_name(slide, "Текст 4"),
        "TalkNFace — «Арена переговоров»\nКейс ЛЦТ 2026 · постановщик ОЭЗ «Алабуга»",
    )
    insert_into_placeholder(slide, logo, "Рисунок 3")


def fill_team(slide, team_logo: Path) -> None:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("Капитан:"):
            put(
                sh,
                "Капитан: Бастанов Айрат Иршатович\n"
                "Специальность: автоматизация тестирования, машинное обучение, прикладное ПО\n"
                "Кол-во участников: 2 человека\n"
                "Как образовалась команда: давно знакомы друг с другом\n"
                "Место учёбы: КФУ\n"
                "Город и регион: не указан\n"
                "Название команды: Ходоки",
                INK,
            )
        elif t.startswith("О команде"):
            put(sh, "О команде", PURPLE)
        elif t.startswith("Краткое описание решения"):
            put(sh, "Краткое описание решения:", PURPLE)
        elif "суть вашего решения" in t:
            put(
                sh,
                "Отрепетируйте сложный разговор до реальной встречи. "
                "Администратор настраивает ситуацию, игрок выбирает ходы, "
                "видит последствия и повторяет ту же версию иначе.",
                INK,
            )
        elif t.startswith("Уникальность решения"):
            put(sh, "Уникальность решения:", PURPLE)
        elif "уникальным или инновационным" in t:
            put(
                sh,
                "Разбор ссылается на сохранённый ход. Отказ не обнуляет процесс. "
                "В показанном продукте нет живого AI: выбор действий считает учебный движок.",
                INK,
            )
    ph = shape_by_name(slide, "Рисунок 1")
    add_picture_fit(slide, team_logo, ph.left, ph.top, ph.width, ph.height)
    drop_shape(ph)


def fill_members(slide) -> None:
    for sh in slide.shapes:
        if sh.has_text_frame and "Заголовок" in sh.name:
            put(sh, "Команда «Ходоки»")
    names = []
    roles = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.startswith("Имя"):
            names.append(sh)
        elif t.startswith("Роль"):
            roles.append(sh)
    names.sort(key=lambda s: int(s.left))
    roles.sort(key=lambda s: int(s.left))
    for i in range(2):
        names[i].height = Inches(0.95)
        roles[i].top = names[i].top + names[i].height + Inches(0.04)
        roles[i].height = Inches(1.55)
    put(names[0], "Бастанов Айрат Иршатович", PURPLE)
    put(
        roles[0],
        "Капитан\nКФУ\nАвтоматизация тестирования, машинное обучение, прикладное ПО\nМессенджер и телефон не переданы",
        INK,
    )
    put(names[1], "Садыков Булат Фаридович", PURPLE)
    put(
        roles[1],
        "Участник\nКФУ\nАвтоматизация тестирования, машинное обучение, прикладное ПО\nМессенджер и телефон не переданы",
        INK,
    )
    for sh in list(slide.shapes):
        if int(sh.left) > int(Inches(5.2)) and int(sh.left) < int(Inches(12.0)):
            drop_shape(sh)


def fill_problem(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "ЗАЧЕМ ТРЕНАЖЁР")
    put(by["Текст 14"], "01", PINK)
    put(by["Текст 1"], "Приём ≠ выбор", PURPLE)
    put(by["Текст 2"], "Знать интересы и пакет недостаточно, чтобы выбрать вопрос или уступку под давлением.", INK)
    put(by["Текст 15"], "02", PINK)
    put(by["Текст 3"], "Лекция", PURPLE)
    put(by["Текст 4"], "Не создаёт этот выбор в момент сделки.", INK)
    put(by["Текст 16"], "03", PINK)
    put(by["Текст 5"], "Живой тренинг", PURPLE)
    put(by["Текст 6"], "Трудно повторить в том же составе, с теми же ставками.", INK)
    put(by["Текст 17"], "04", PINK)
    put(by["Текст 7"], "Игрок", PURPLE)
    put(by["Текст 8"], "Начинающий закупщик или руководитель. Ошибка без цены реальной сделки.", INK)
    put(by["Текст 18"], "05", PINK)
    put(by["Текст 9"], "Администратор", PURPLE)
    put(by["Текст 10"], "Задаёт шесть полей и публикует ситуацию. Чужие попытки и отчёты ему закрыты.", INK)
    put(by["Текст 19"], "06", PINK)
    put(by["Текст 11"], "Наставник", PURPLE)
    put(by["Текст 12"], "Разбирает отчёт на экране игрока. Отдельного доступа к чужим попыткам нет.", INK)


def fill_journey(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "КАК ПРОХОДИТ ВСТРЕЧА", PINK)
    put(by["Текст 14"], "01", PINK)
    put(by["Текст 1"], "Настроить", PURPLE)
    put(
        by["Текст 2"],
        "Шесть полей: сфера, тема, сложность, тон, роль, цель. "
        "Недостижимая цель не публикуется. До 8 ходов; встреча может закончиться раньше.",
        INK,
    )
    put(by["Текст 15"], "02", PINK)
    put(by["Текст 3"], "Выбрать", PURPLE)
    put(
        by["Текст 4"],
        "Игрок без регистрации фиксирует цель и альтернативу, затем выбирает ход: "
        "вопрос, признание факта, довод, пакет, принятие, выход.",
        INK,
    )
    put(by["Текст 16"], "03", PINK)
    put(by["Текст 5"], "Разобрать", PURPLE)
    put(
        by["Текст 6"],
        "Исход сделки отдельно от процесса. Ссылка открывает сохранённый ход. "
        "«Повторить ту же ситуацию» создаёт новую попытку той же версии.",
        INK,
    )


def fill_settings(slide, settings_png: Path) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 2"], "Шесть полей администратора", PURPLE)
    put(
        by["Текст 6"],
        "Сфера · тема · сложность · тон · роль · цель.\n"
        "Сложность «обычная» держит до 8 ходов.\n"
        "Недостижимая цель блокируется до публикации.\n"
        "Справа — форма настройки. Это не живой показ "
        "«плановая / дружелюбный / маржа»: проверенная пара исходов на следующем слайде "
        "идёт по неизменённому эталону «Закупка к запуску».",
        INK,
    )
    drop_shape(shape_by_name(slide, "Рисунок 3"))
    slide.shapes.add_picture(str(settings_png), Inches(6.9), Inches(1.2), Inches(6.0), Inches(5.7))


def fill_families(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "ДВЕ СЕМЬИ СЦЕНАРИЕВ")
    put(by["Текст 14"], "S1", PINK)
    put(by["Текст 1"], "Закупка к запуску", PURPLE)
    put(by["Текст 2"], "Промышленные закупки. Цена партии, график поставки, предоплата.", INK)
    put(by["Текст 15"], "S2", PINK)
    put(by["Текст 3"], "Срочная задача", PURPLE)
    put(by["Текст 4"], "Нагрузка команды. Объём, срок, помощник, перенос отчёта.", INK)
    put(by["Текст 16"], "03", PINK)
    put(by["Текст 5"], "Граница", PURPLE)
    put(by["Текст 6"], "Третьей отрасли нет. Не каждая комбинация шести полей проходит проверку.", INK)
    put(by["Текст 17"], "04", PINK)
    put(by["Текст 7"], "Баллы", PURPLE)
    put(
        by["Текст 8"],
        "Полезность описывает пакет в учебной модели. "
        "100 процесса — применимые проверки записи, не компетенция человека.",
        INK,
    )


def fill_attempts(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "ДВА ИСХОДА ОДНОГО ЭТАЛОНА")
    put(by["Текст 14"], "Попытка А", PURPLE)
    put(by["Текст 16"], "Попытка Б", PURPLE)
    put(
        by["Текст 2"],
        "Неизменённый эталон «Закупка к запуску»: обычная сложность, нейтральный тон, "
        "директор продаж, оборотные средства.\n"
        "Вопрос про логистику → признание экономии на разделении партии → довод → "
        "вопрос про оплату → пакет 95 / 40% на день 7 и остаток на 14 / предоплата 50%.\n"
        "Исход: взаимная выгода. Учебная полезность 64. Процесс 100.",
        INK,
    )
    put(
        by["Текст 6"],
        "Тот же эталон. Сначала тот же пакет, затем вопросы, затем явный выход из переговоров.\n"
        "Исход: нет соглашения. Полезность сделки не считается. Процесс 70. "
        "Зачёт взаимно приемлемого пакета сохраняется.\n"
        "Отказ — отдельный ход игрока, а не следствие порядка вопросов.",
        INK,
    )


def fill_evidence(slide, evidence_png: Path) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "РАЗБОР ПО ХОДАМ")
    put(by["Текст 3"], "Проверка «выполнено» со ссылкой «Ход 5 — открыть сохранённый момент».", INK)
    put(by["Текст 4"], "Сохранённая фраза пакета: 95 / 40% на день 7, остаток на 14 / предоплата 50%.", INK)
    put(by["Текст 5"], "Если наблюдений мало или нет явной подготовки, общий балл процесса не выводится.", INK)
    ph = shape_by_name(slide, "Рисунок 1")
    add_picture_fit(slide, evidence_png, ph.left, ph.top, ph.width, ph.height)
    drop_shape(ph)


def fill_quiz(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 27"], "Это не тест на 100 баллов", PURPLE)
    put(
        by["Объект 28"],
        "100 из 100 — применимые проверки этой записи, не оценка компетенции человека.\n"
        "Отказ не равен провалу процесса: попытка Б без соглашения получила процесс 70.\n"
        "Слабое соглашение в семье нагрузки: учебная полезность 20, процесс 44,12.\n"
        "Администратор не читает чужие отчёты. Наставник обсуждает разбор на экране игрока.",
        INK,
    )


def fill_marketing(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "КОМУ И КАКАЯ ГИПОТЕЗА")
    put(by["Текст 14"], "Маркетинг", PURPLE)
    put(by["Текст 16"], "Бизнес", PURPLE)
    put(
        by["Текст 2"],
        "Целевой сегмент: HR / L&D, начинающий закупщик и руководитель в обучении.\n"
        "Пилотов, договоров и названных клиентов нет.\n"
        "ОЭЗ «Алабуга» — постановщик кейса хакатона, не клиент продукта.",
        INK,
    )
    put(
        by["Текст 6"],
        "Гипотеза пользы: безопасные попытки дешевле ошибки на реальной встрече.\n"
        "Измеренного ROI, выручки и воронки нет — цифры не выдуманы.\n"
        "Файла LICENSE в репозитории нет. Это решение владельца, не дисквалификация из PDF кейса.",
        INK,
    )


def fill_tech(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "КАК УСТРОЕНО")
    put(
        by["Текст 2"],
        "Браузер. Один origin: Fastify + React. Показ слушает 127.0.0.1. Регистрации игрока нет.",
        INK,
    )
    put(
        by["Текст 3"],
        "Учебный движок считает допустимость, раскрытие фактов и полезность. "
        "Модель, CUDA и платный API для игры не нужны.",
        INK,
    )
    put(
        by["Текст 4"],
        "SQLite хранит публикации и попытки. Администратор не обходит границу чужих отчётов.",
        INK,
    )


def fill_next(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 13"], "СЕЙЧАС И ДАЛЬШЕ")
    put(by["Текст 14"], "Работает сейчас", PURPLE)
    put(by["Текст 16"], "Не в этом демо", PURPLE)
    put(
        by["Текст 2"],
        "Выбор действий, две семьи, шесть полей, последствия, разбор, повторная попытка.\n"
        "Локальный запуск на Windows без модели и ключа API.\n"
        "Команды Prepare и Start — отдельно от Stop. Адрес после запуска: http://127.0.0.1:3100/admin на том компьютере.",
        INK,
    )
    put(
        by["Текст 6"],
        "Свободный текст, AI-оппонент, генерация сценария моделью.\n"
        "Qwen не встроена и по качеству не проверена.\n"
        "Публичного сайта нет. 127.0.0.1 не является ссылкой для формы хакатона.",
        INK,
    )


def add_link_para(tf, label: str, url: str) -> None:
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = label
    run.font.color.rgb = PINK
    run.hyperlink.address = url


def fill_materials(slide) -> None:
    by = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}
    put(by["Заголовок 1"], "Материалы в репозитории")
    body = shape_by_name(slide, "Объект 2")
    put(
        body,
        "Команда «Ходоки», 2 человека, КФУ. Публичного сайта нет: прототип — локальный запуск.",
        INK,
    )
    tf = body.text_frame
    add_link_para(tf, "Документация — docs/submission/README.md", BLOB + "docs/submission/README.md")
    add_link_para(
        tf,
        "Презентация — docs/submission/TALKNFACE_PRODUCT_PITCH_RU.pdf",
        BLOB + "docs/submission/TALKNFACE_PRODUCT_PITCH_RU.pdf",
    )
    add_link_para(tf, "Прототип — docs/DEMO_GUIDE_RU.md", BLOB + "docs/DEMO_GUIDE_RU.md")
    add_link_para(
        tf,
        "Дополнительно — docs/submission/PRODUCT_OVERVIEW_RU.md",
        BLOB + "docs/submission/PRODUCT_OVERVIEW_RU.md",
    )
    add_link_para(
        tf,
        "PPTX и адреса коммита — docs/submission/FORM_LINKS_RU.md",
        BLOB + "docs/submission/FORM_LINKS_RU.md",
    )


def build() -> None:
    template = find_template()
    if template.resolve() == PPTX.resolve():
        raise RuntimeError("Refusing to treat the submission PPTX as the official template.")
    assets = prepare_assets()
    logo = extract_alabuga(template)
    assemble(template)
    prs = Presentation(str(ASSEMBLED))
    if len(prs.slides) != 14:
        raise RuntimeError(f"Expected 14 slides, got {len(prs.slides)}")
    slides = list(prs.slides)
    fill_title(slides[0], logo)
    fill_team(slides[1], assets["khodoki"])
    fill_members(slides[2])
    fill_problem(slides[3])
    fill_journey(slides[4])
    fill_settings(slides[5], assets["settings"])
    fill_families(slides[6])
    fill_attempts(slides[7])
    fill_evidence(slides[8], assets["evidence"])
    fill_quiz(slides[9])
    fill_marketing(slides[10])
    fill_tech(slides[11])
    fill_next(slides[12])
    fill_materials(slides[13])
    strip_notes(prs)
    PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(PPTX))
    print(f"wrote {PPTX} slides={len(prs.slides)}")


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"build failed: {exc}", file=sys.stderr)
        raise
